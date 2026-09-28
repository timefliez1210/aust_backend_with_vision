#!/usr/bin/env python3
"""
Merge Alex's reviewed Rechnungsausgangsbuch into the invoices table.

Reads his book, his answers in the Abgleich workbook, and the CURRENT state of the
target database, and writes ONE reviewable SQL file that does the whole merge inside a
single transaction:

    python3 scripts/register-import/import_register.py \\
        --book "Rechnungsausgangsbuch 2024.xlsx" --sheet 2026 \\
        --review Abgleich_2026.xlsx \\
        --psql "docker exec -i aust_staging_postgres psql -U aust_staging -d aust_import_test" \\
        --out import_2026.sql [--apply]

Nothing is written unless `--apply` is given; without it the script only prints the
plan and writes the SQL file. `--apply` pipes that same file into `--psql`.

# Why the SQL re-checks everything
The file is generated from a snapshot and may be applied minutes (or a restore) later.
Every prod row it touches is pinned by a fingerprint of its current contents, and every
number it inserts is asserted free at insert time. If anything moved in between —
Alex booked a payment, a new invoice took a number — the transaction aborts and
nothing is written. Regenerate and re-run.

# Rules (decided with the user, see memory project_register_excel_import_2026_08_23)
- His book is authoritative for numbers, amounts, dates, Zahlungsart and Bemerkungen.
- The seven never-sent prod drafts sitting on his numbers 01–09 ARE invoices he later
  wrote by hand under other numbers; they are moved onto those numbers (RENUMBER),
  not duplicated.
- A payment prod recorded after his 20.08. snapshot is never undone.
- An unreadable Bezahlt-Datum is never guessed: the money is booked as received in
  full (Teilzahlung = Brutto), the date stays empty, the raw text goes into the note.
- Ledger-only rows carry the customer directly (`customer_id`), no inquiry, no PDF,
  `is_legacy = TRUE`, and their Leistungszeitraum in `service_start/_end`.
"""

import argparse
import datetime as dt
import json
import re
import subprocess
import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_review_workbook import as_date, raw_date_text, read_excel, tokens  # noqa: E402

import openpyxl  # noqa: E402

YEAR = 2026
TODAY = dt.date.today().strftime("%d.%m.%Y")

# ── Reconciliation decisions ────────────────────────────────────────────────
#
# Prod number → the number the same invoice carries in Alex's book. Every entry here
# was a draft that never left the building except 0059 and 0067, which Alex asked to
# move himself (book notes on 2026-35 and 2026-67, review note on 2026-68).
RENUMBER = {
    "2026-0003": "2026-20",   # Daniel Krause 624,00
    "2026-0002": "2026-22",   # Jonny Heidemann 3.506,00
    "2026-0001": "2026-23",   # Bauche — book amount differs, see AMOUNT_FROM_BOOK
    "2026-0004": "2026-27",   # Anzahlung Lullies 378,00 (30 % of 1.260)
    "2026-0005": "2026-43",   # Restzahlung Lullies 882,00
    "2026-0008": "2026-44",   # Anzahlung Eggert — book split differs
    "2026-0009": "2026-67",   # Restrechnung Eggert — book split differs
    "2026-0059": "2026-35",   # Gutschrift Stache: "muss online geändert werden von 59 auf RG 35"
    "2026-0067": "2026-67A",  # Lippmann: "Lippman zu 2026-0067A machen"
}

# Numbers (book) whose prod amount may be overwritten by the book amount. Anything else
# that disagrees aborts the run: a silent amount change in a legal register is the one
# mistake this importer must never make.
AMOUNT_FROM_BOOK = {
    23,  # Bauche draft 1.316,35 → book 1.153,35 (he billed 23 + 39 by hand)
    44,  # Eggert Anzahlung draft 533,40 → book 620,40
    67,  # Eggert Rest draft 1.244,60 → book 1.157,60 (sum 1.778 unchanged)
    66,  # Lalicata: stored base 400,00, offer and book 640,00
    50,  # ifm Anzahlung: invoice 1.478,80, book 1.478,00 — his book is the source of truth
}

# Book numbers whose prod invoice goes to a different customer: Alex issued it himself,
# to the person his book names (2026-70 was sent by him to Kampe on 03.08.; prod's
# never-sent draft sat on Lier's inquiry).
CUSTOMER_FROM_BOOK = {70: "Gabriele Kampe"}

# Bezahlt-Datum cells that are not a date, resolved with the user (2026-09-28).
PAID_DATE_OVERRIDE = {
    17: dt.date(2026, 6, 30),  # book: "31.06.26"
}

# Book customer → prod customer, where the book names the company and prod holds the
# contact person. Evidence: prod already carries the company in `company_name`, and the
# invoices with the same number and amount sit on that contact (review rows 56/63/80/86).
CUSTOMER_ALIAS = {
    "erfi Ernst Fischer GmbH+Co.KG": "019e4190-4b7a-7272-be3e-4ca8046d90a6",  # Finkbeiner (review note 2026-79)
    "Hildesheimer Dienste": "019f18ec-7316-7490-8b39-ce940cc7e785",           # Nicolai Bocancea
    "Transcome GmbH": "019d680e-e0ac-7a60-a216-5f8f7bb50627",                 # Stephan Schrader
    "Blankenstein Logistik GmbH": "019d3b96-b23c-7721-9735-cab64fa5f5a4",     # Martin Blankenstein
    "Luttert Ordnungs u.Regal Systeme": "019dbedf-41d8-7bb0-bfb0-47e8c463be5f",
    "Praxis Günter Engelhardt": "019f17b0-994f-7de0-9040-f23a2ac3bfc8",
    "Engelhardt": "019f17b0-994f-7de0-9040-f23a2ac3bfc8",                     # "LagerungEngelhardt"
    "Lingolf Hermann, LL.M.": "019cfb35-410f-75b2-86a5-2e341447b36a",
    "Lingolf Hermann": "019cfb35-410f-75b2-86a5-2e341447b36a",
    "Bettels": "01a05c37-9cd3-7920-85a8-e843c8ebff9a",                        # not the test2@test.com twin
    "Michael Stache": "019d1ac7-e68b-7401-a5ca-2907edb3526d",
    "Monika Stache": "019d1ac7-e68b-7401-a5ca-2907edb3526d",
}

# Customers that exist nowhere in prod. `customer_type` follows the legal form.
NEW_CUSTOMERS = {
    "AllTransport GmbH": "business",
    "Feelings Braut- & Festmoden": "business",
    "Haßenpflug": "private",
    "Löwen Dienstleistungen & Transport": "business",
    "Steinberg GmbH": "business",
    "Gabriele Kampe": "private",
}

# Review notes we understood and encoded above. Any other note aborts the run.
KNOWN_REVIEW_NOTES = {
    "2026-68": "67a",
    "2026-79": "Finkbeiner is der Ansprechpartner der Firma",
}

TYPE_LABEL = {
    "partial_first": "Anzahlung",
    "partial_final": "Restzahlung",
    "gutschrift": "Gutschrift",
    "lagerung": "Lagerung",
}


# ── Helpers ─────────────────────────────────────────────────────────────────

def seq_of(number: str):
    m = re.fullmatch(r"(\d{4})-0*(\d+)", number.strip())
    return (int(m.group(1)), int(m.group(2))) if m else None


def q(v) -> str:
    """SQL literal."""
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, dt.date):
        return f"'{v.isoformat()}'"
    return "'" + str(v).replace("'", "''") + "'"


def ts(d):
    """A book date as a timestamp: noon Berlin, so no timezone can move the day."""
    return "NULL" if d is None else f"('{d.isoformat()} 12:00'::timestamp AT TIME ZONE 'Europe/Berlin')"


def eur(cents: int) -> str:
    return f"{cents / 100:,.2f} €".translate(str.maketrans({",": ".", ".": ","}))


def brutto(netto: int) -> int:
    return round(netto * 1.19)


def parse_period(v):
    """Alex's Datum column → (start, end, raw_if_unreadable).

    Forms in his book: a real date; '12-13.01.2026'; '07.11.25-30.01.2026';
    '23.01-23.02.2026'; '24.03.-30.09.26'; '23.-24.07.26'.
    """
    d = as_date(v)
    if d:
        return d, None, ""
    if v is None:
        return None, None, ""
    raw = str(v).strip()

    def full(s):
        parts = [p for p in s.split(".") if p]
        if len(parts) != 3:
            return None
        day, month, year = (int(p) for p in parts)
        if year < 100:
            year += 2000
        try:
            return dt.date(year, month, day)
        except ValueError:
            return None

    try:
        if "-" not in raw:
            d = full(raw)
            return (d, None, "") if d else (None, None, raw)
        left, right = raw.split("-", 1)
        end = full(right)
        if not end:
            return None, None, raw
        lp = [int(p) for p in left.split(".") if p]
        if len(lp) == 1:
            start = dt.date(end.year, end.month, lp[0])
        elif len(lp) == 2:
            start = dt.date(end.year if lp[1] <= end.month else end.year - 1, lp[1], lp[0])
        elif len(lp) == 3:
            start = dt.date(lp[2] + 2000 if lp[2] < 100 else lp[2], lp[1], lp[0])
        else:
            return None, None, raw
        if start > end:
            return None, None, raw
        return start, (end if end != start else None), ""
    except ValueError:
        return None, None, raw


def lenient_date(v):
    """A date cell, or one whose only fault is a missing dot ('31.0326', '29.062026').

    Those two are unambiguous. Anything else — '31.06.26' is a day June does not
    have — stays unreadable; the importer never picks a date for Alex.
    """
    d = as_date(v)
    if d or v is None:
        return d
    raw = str(v).strip()
    # Only the exact shapes "DD.MMYY" and "DD.MMYYYY": one dot missing, nothing else wrong.
    m = re.fullmatch(r"(\d{2})\.(\d{2})(\d{2}|\d{4})", raw)
    if not m:
        return None
    year = int(m.group(3)) + (2000 if len(m.group(3)) == 2 else 0)
    try:
        return dt.date(year, int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def compute_netto(p) -> int:
    """Mirror of routes/invoices.rs::compute_invoice_amounts (netto only)."""
    base = p["base_netto_cents"] if p["base_netto_cents"] is not None else (p["offer_netto_cents"] or 0)
    extras = sum(int(e.get("price_cents", 0)) for e in (p["extra_services"] or []))
    if p["is_manual"]:
        return round(sum(float(li["quantity"]) * int(li["unit_price_cents"]) for li in (p["line_items_json"] or [])))
    pct = p["partial_percent"] if p["partial_percent"] is not None else (p["deposit_percent"] or 0)
    offer_brutto = round(base * 1.19)
    first_brutto = round(offer_brutto * pct / 100)
    first_netto = round(first_brutto / 1.19)
    if p["invoice_type"] == "partial_first":
        return first_netto
    if p["invoice_type"] == "partial_final":
        return base - first_netto + extras
    return base + extras


def run_psql(psql, sql):
    out = subprocess.run(psql + " -X -q -tA -v ON_ERROR_STOP=1", shell=True, input=sql,
                         capture_output=True, text=True)
    if out.returncode != 0:
        sys.exit(f"psql failed:\n{out.stderr}")
    return out.stdout


FINGERPRINT = ("md5(row(inv.invoice_number, inv.status, inv.invoice_type, inv.sent_at, inv.paid_at, "
               "inv.base_netto_cents, inv.extra_services, inv.line_items_json, inv.notes, "
               "inv.payment_method, inv.customer_id, inv.paid_amount_cents)::text)")


def load_db(psql):
    sql = f"""
    SELECT json_build_object(
      'invoices', (SELECT coalesce(json_agg(x), '[]') FROM (
        SELECT inv.id, inv.invoice_number, inv.invoice_type, inv.partial_percent, inv.deposit_percent,
               inv.status, inv.is_manual, inv.extra_services, inv.line_items_json, inv.base_netto_cents,
               inv.sent_at::date AS sent, inv.paid_at::date AS paid, inv.payment_method, inv.notes,
               inv.due_date, inv.paid_amount_cents, inv.pdf_s3_key, inv.is_legacy,
               c.name AS customer, c.id AS customer_id, {FINGERPRINT} AS fp,
               (SELECT o.price_cents FROM offers o WHERE o.inquiry_id = inv.inquiry_id
                  AND o.status NOT IN ('rejected','cancelled','superseded')
                ORDER BY o.created_at DESC LIMIT 1) AS offer_netto_cents
        FROM invoices inv
        LEFT JOIN inquiries i ON i.id = inv.inquiry_id
        LEFT JOIN customers c ON c.id = COALESCE(inv.customer_id, i.customer_id)
        WHERE inv.invoice_number LIKE '{YEAR}-%') x),
      'storage', (SELECT coalesce(json_agg(invoice_number), '[]') FROM storage_invoices
                  WHERE invoice_number LIKE '{YEAR}-%'),
      'customers', (SELECT coalesce(json_agg(json_build_object('id', id, 'name', name)), '[]')
                    FROM customers WHERE merged_into IS NULL),
      'has_service_cols', EXISTS (SELECT 1 FROM information_schema.columns
                    WHERE table_name = 'invoices' AND column_name = 'service_start')
    )"""
    return json.loads(run_psql(psql, sql))


def read_review(path):
    wb = openpyxl.load_workbook(path)
    problems = []
    for r in wb["Abgleich"].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        if r[9] != "Passt so":
            problems.append(f"{r[0]}: Entscheidung „{r[9]}“ ist nicht umgesetzt")
        if r[10] and KNOWN_REVIEW_NOTES.get(r[0]) != r[10]:
            problems.append(f"{r[0]}: unbekannte Notiz „{r[10]}“")
    for r in wb["Kunden"].iter_rows(min_row=2, values_only=True):
        if r[0] and r[3] != "Neu anlegen":
            problems.append(f"Kunde {r[0]}: Entscheidung „{r[3]}“ ist nicht umgesetzt")
    if problems:
        sys.exit("Review enthält Antworten, die dieses Skript nicht kennt:\n  " + "\n  ".join(problems))


# ── Plan ────────────────────────────────────────────────────────────────────

def plan(book, db):
    if not db["has_service_cols"]:
        sys.exit("invoices.service_start fehlt — Migration 20260927120000 zuerst anwenden.")
    if db["storage"]:
        sys.exit(f"storage_invoices hat {YEAR}-Nummern ({db['storage']}) — nicht vorgesehen, bitte prüfen.")

    by_number = {p["invoice_number"]: p for p in db["invoices"]}
    for old in RENUMBER:
        if old not in by_number:
            sys.exit(f"{old} existiert nicht mehr — Plan passt nicht zur Datenbank.")

    # Where every prod invoice sits after the renumbering, keyed by book sequence.
    placed = {}
    for p in db["invoices"]:
        number = RENUMBER.get(p["invoice_number"], p["invoice_number"])
        s = seq_of(number)
        if s and s[0] == YEAR:
            if s[1] in placed:
                sys.exit(f"Nummer {YEAR}-{s[1]} wäre doppelt belegt ({placed[s[1]]['invoice_number']}, {p['invoice_number']}).")
            placed[s[1]] = p

    customers = {c["id"]: c["name"] for c in db["customers"]}
    existing_names = {(c["name"] or "").strip().lower() for c in db["customers"]}
    for name in NEW_CUSTOMERS:
        if name.lower() in existing_names:
            sys.exit(f"Kunde „{name}“ existiert inzwischen — NEW_CUSTOMERS anpassen.")
    for alias, cid in CUSTOMER_ALIAS.items():
        if cid not in customers:
            sys.exit(f"Alias-Kunde {cid} ({alias}) fehlt.")

    def customer_for(book_name):
        name = book_name.strip()
        if name in CUSTOMER_ALIAS:
            return CUSTOMER_ALIAS[name], None
        if name in NEW_CUSTOMERS:
            return None, name
        want = tokens(name)
        exact = [cid for cid, cname in customers.items() if cname and tokens(cname) == want]
        wider = [cid for cid, cname in customers.items() if cname and want and want <= tokens(cname)]
        for hits in (exact, wider):
            if len(hits) == 1:
                return hits[0], None
            if len(hits) > 1:
                break
        sys.exit(f"Kunde „{name}“ ist nicht eindeutig zuzuordnen ({len(hits)} Treffer) — CUSTOMER_ALIAS ergänzen.")

    updates, inserts, report = [], [], []
    seen = set()
    for b in book:
        n = b["nr"]
        seen.add(n)
        start, end, raw_period = parse_period(b["leistung"])
        sent = lenient_date(b["versendet"])
        paid = PAID_DATE_OVERRIDE.get(n) or lenient_date(b["bezahlt"])
        due = as_date(b["faellig"])
        says_paid = str(b["offen"]).strip().lower() == "bezahlt"
        notes = [b["bemerkung"]] if b["bemerkung"] else []
        unreadable_paid = says_paid and paid is None and raw_date_text(b["bezahlt"])
        if n in PAID_DATE_OVERRIDE:
            notes.append(f"Bezahlt am lt. Buch: „{raw_date_text(b['bezahlt'])}“ → "
                         f"{paid:%d.%m.%Y} gebucht (Absprache 28.09.2026)")
        if unreadable_paid:
            notes.append(f"Bezahlt am lt. Buch: „{raw_date_text(b['bezahlt'])}“ (unlesbar)")
        if b["versendet"] is not None and sent is None:
            notes.append(f"Versendet lt. Buch: „{b['versendet']}“ (unlesbar)")
        if raw_period:
            notes.append(f"Leistung lt. Buch: „{raw_period}“")

        p = placed.get(n)
        if p is None:
            cid, new_name = customer_for(b["kunde"])
            label = TYPE_LABEL.get(b["typ"])
            if label or b["raw_name"].strip() != (customers.get(cid) or new_name):
                notes.insert(0, f"Buch: {b['raw_name'].strip()}")
            status = "paid" if paid else "sent"
            inserts.append(dict(
                number=f"{YEAR}-{n:02d}", customer_id=cid, new_customer=new_name,
                netto=b["netto"], sent=sent, paid=paid, due=due, status=status,
                method=b["art"] or None, notes=" · ".join(notes) or None,
                start=start, end=end,
                paid_amount=brutto(b["netto"]) if unreadable_paid else None,
                created=sent or paid or start,
            ))
            report.append(f"NEU      {YEAR}-{n:02d}  {b['raw_name'].strip()[:40]:40} {eur(b['netto']):>12}  → {customers.get(cid) or new_name + ' (neu)'}")
            continue

        # Existing prod invoice: book values win, except a payment recorded after his snapshot.
        set_ = {}
        audit = []
        old_number = p["invoice_number"]
        new_number = RENUMBER.get(old_number)
        if new_number:
            set_["invoice_number"] = q(new_number)
            audit.append(f"vorher Nr. {old_number}")

        have = compute_netto(p)
        if abs(have - b["netto"]) > 2:
            if n not in AMOUNT_FROM_BOOK:
                sys.exit(f"{YEAR}-{n:02d}: Betrag System {eur(have)} ≠ Buch {eur(b['netto'])} und nicht freigegeben.")
            audit.append(f"vorher {eur(have)} netto")
            if p["invoice_type"] != "full":
                audit.append(f"vorher {TYPE_LABEL[p['invoice_type']]}"
                             + (f" {p['partial_percent']} %" if p["partial_percent"] else ""))
            if p["is_manual"]:
                sys.exit(f"{YEAR}-{n:02d}: manuelle Rechnung, Betrag kann nicht umgeschrieben werden.")
            set_.update(invoice_type="'full'", partial_percent="NULL", partial_group_id="NULL",
                        extra_services="'[]'::jsonb", base_netto_cents=str(b["netto"]))

        was_unsent_draft = p["sent"] is None and p["paid"] is None
        if new_number and was_unsent_draft:
            # The draft PDF shows the old number and, for three of them, the wrong amount.
            set_["pdf_s3_key"] = "NULL"
        elif new_number and p["pdf_s3_key"]:
            audit.append(f"PDF zeigt noch {old_number}")

        if paid:
            set_["paid_at"] = ts(paid)
            set_["status"] = "'paid'"
        elif unreadable_paid and p["paid"] is None:
            set_["paid_amount_cents"] = str(brutto(b["netto"]))
            set_["status"] = "'sent'"
        # else: keep whatever prod has — it may know about a later payment.

        if sent:
            set_["sent_at"] = ts(sent)
            if "status" not in set_ and p["status"] in ("draft", "ready"):
                set_["status"] = "'sent'"
        elif paid and (p["sent"] is None or p["sent"] == p["paid"]):
            # No Rechnungsdatum in the book; prod's is the mark-paid backfill (same day as
            # the booking). Re-derive it from the real payment date, as mark_paid would.
            set_["sent_at"] = ts(paid)
        elif was_unsent_draft and says_paid:
            set_["status"] = "'sent'"
        if due:
            set_["due_date"] = q(due)
        if b["art"] and b["art"] != p["payment_method"]:
            set_["payment_method"] = q(b["art"])

        if n in CUSTOMER_FROM_BOOK:
            set_["customer_id"] = ("NEW", CUSTOMER_FROM_BOOK[n])
            audit.append(f"vorher Kunde {p['customer']}")
        # The book names someone prod does not recognise. The invoice keeps its
        # customer, but Alex's name must not vanish from the record.
        elif (CUSTOMER_ALIAS.get(b["kunde"].strip()) != p["customer_id"]
                and not (tokens(b["kunde"]) & tokens(p["customer"] or ""))):
            notes.insert(0, f"Buch: {b['raw_name'].strip()}")

        old_notes = p["notes"] or ""
        add = [x for x in notes if x.strip() and x.strip() not in old_notes]
        if audit:
            add.append(f"Import {TODAY}: " + ", ".join(audit))
        if add:
            set_["notes"] = q(" · ".join([old_notes] + add if old_notes else add))

        if set_:
            updates.append(dict(id=p["id"], fp=p["fp"], number=old_number, set=set_))
        report.append(
            f"{'UMNUMM.' if new_number else 'ABGLEICH'} {old_number:>9}{' → ' + new_number if new_number else '':12} "
            f"{p['customer'] or '':24.24} {eur(have):>12}"
            + (f" → {eur(b['netto'])}" if abs(have - b['netto']) > 2 else "")
            + f"  [{', '.join(k for k in set_)}]"
        )

    # Prod rows the book does not mention stay as they are.
    for s, p in sorted(placed.items()):
        if s not in seen:
            report.append(f"BLEIBT   {p['invoice_number']:>9}{' → ' + RENUMBER.get(p['invoice_number'], ''):12} {p['customer'] or ''} (nicht im Buch)")
    # Renumbers onto a number outside the book's sequence (67A) have no book row to
    # carry them; they are moved as they are.
    for p in db["invoices"]:
        target = RENUMBER.get(p["invoice_number"])
        if target and not seq_of(target):
            old_notes = p["notes"] or ""
            audit = f"Import {TODAY}: vorher Nr. {p['invoice_number']}" + (
                f", PDF zeigt noch {p['invoice_number']}" if p["pdf_s3_key"] else "")
            updates.append(dict(id=p["id"], fp=p["fp"], number=p["invoice_number"], set={
                "invoice_number": q(target),
                "notes": q(f"{old_notes} · {audit}" if old_notes else audit),
            }))
            report.append(f"UMNUMM. {p['invoice_number']:>9} → {target:9} {p['customer']} (nicht im Buch, Alex: „67A“)")

    return updates, inserts, report


# ── SQL ─────────────────────────────────────────────────────────────────────

def render_sql(updates, inserts, book_total_netto):
    out = [
        f"-- Register import {YEAR}, generated {dt.datetime.now():%Y-%m-%d %H:%M}.",
        "-- One transaction: any failed check rolls back everything.",
        "\\set ON_ERROR_STOP on",
        "BEGIN;",
        "",
        "-- 1. Every prod row we touch must still be exactly as planned.",
    ]
    for u in updates:
        out.append(
            f"DO $$ BEGIN IF NOT EXISTS (SELECT 1 FROM invoices inv WHERE inv.id = {q(u['id'])} "
            f"AND {FINGERPRINT} = {q(u['fp'])}) THEN RAISE EXCEPTION 'Rechnung {u['number']} hat sich seit der Planung geändert'; END IF; END $$;"
        )

    out += ["", "-- 2. New customers."]
    new_ids = {}
    wanted = [ins["new_customer"] for ins in inserts] + [
        v[1] for u in updates for v in u["set"].values() if isinstance(v, tuple)]
    for name in wanted:
        if name and name not in new_ids:
            new_ids[name] = str(uuid.uuid7())
            ctype = NEW_CUSTOMERS[name]
            out.append(
                f"INSERT INTO customers (id, name, customer_type, company_name, notes) VALUES "
                f"({q(new_ids[name])}, {q(name)}, {q(ctype)}, {q(name if ctype == 'business' else None)}, "
                f"{q(f'Angelegt beim Import des Rechnungsausgangsbuchs {YEAR} ({TODAY})')});"
            )

    out += ["", "-- 3. Renumber and reconcile existing invoices (renumbers first, so their old numbers are free)."]
    for u in sorted(updates, key=lambda u: "invoice_number" not in u["set"]):
        sets = ", ".join(f"{k} = {q(new_ids[v[1]]) if isinstance(v, tuple) else v}"
                         for k, v in u["set"].items())
        out.append(f"UPDATE invoices SET {sets} WHERE id = {q(u['id'])};  -- {u['number']}")

    out += ["", "-- 4. Ledger-only rows from the book."]
    for ins in inserts:
        s = seq_of(ins["number"])[1]
        cid = ins["customer_id"] or new_ids[ins["new_customer"]]
        out.append(
            f"DO $$ BEGIN IF EXISTS (SELECT 1 FROM invoices WHERE invoice_number ~ '^{YEAR}-0*{s}$') "
            f"OR EXISTS (SELECT 1 FROM storage_invoices WHERE invoice_number ~ '^{YEAR}-0*{s}$') "
            f"THEN RAISE EXCEPTION 'Nummer {ins['number']} ist belegt'; END IF; END $$;"
        )
        out.append(
            "INSERT INTO invoices (id, inquiry_id, customer_id, invoice_number, invoice_type, status, "
            "base_netto_cents, sent_at, paid_at, due_date, payment_method, notes, paid_amount_cents, "
            "service_start, service_end, is_legacy, created_at) VALUES ("
            f"{q(str(uuid.uuid7()))}, NULL, {q(cid)}, {q(ins['number'])}, 'full', {q(ins['status'])}, "
            f"{ins['netto']}, {ts(ins['sent'])}, {ts(ins['paid'])}, {q(ins['due'])}, {q(ins['method'])}, "
            f"{q(ins['notes'])}, {q(ins['paid_amount'])}, {q(ins['start'])}, {q(ins['end'])}, TRUE, "
            f"{ts(ins['created']) if ins['created'] else 'now()'});"
        )

    out += [
        "",
        "-- 5. Post-conditions: every book number exactly once, and the counter ahead of all of them.",
        f"""DO $$ DECLARE missing text; dup text; BEGIN
  SELECT string_agg(n::text, ',') INTO missing FROM generate_series(1, {max(i for i in BOOK_SEQS)}) n
   WHERE n IN ({','.join(map(str, sorted(BOOK_SEQS)))})
     AND NOT EXISTS (SELECT 1 FROM invoices WHERE invoice_number ~ ('^{YEAR}-0*' || n || '$'));
  IF missing IS NOT NULL THEN RAISE EXCEPTION 'Fehlende Nummern: %', missing; END IF;
  SELECT string_agg(k, ',') INTO dup FROM (
    SELECT (regexp_match(invoice_number, '^{YEAR}-0*(\\d+)$'))[1] AS k FROM invoices
     WHERE invoice_number ~ '^{YEAR}-\\d+$' GROUP BY 1 HAVING count(*) > 1) d;
  IF dup IS NOT NULL THEN RAISE EXCEPTION 'Doppelte Nummern: %', dup; END IF;
  IF (SELECT max((regexp_match(invoice_number, '^{YEAR}-0*(\\d+)$'))[1]::int) FROM invoices)
     > (SELECT last_value FROM invoice_number_counters WHERE year = {YEAR}) THEN
    RAISE EXCEPTION 'invoice_number_counters liegt hinter der höchsten Nummer';
  END IF;
END $$;""",
        "",
        "COMMIT;",
    ]
    return "\n".join(out) + "\n"


BOOK_SEQS: set = set()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--book", required=True)
    ap.add_argument("--sheet", default=str(YEAR))
    ap.add_argument("--review", required=True)
    ap.add_argument("--psql", required=True, help="command that runs psql against the target DB, reading SQL on stdin")
    ap.add_argument("--out", required=True)
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    read_review(args.review)
    book = read_excel(args.book, args.sheet)
    BOOK_SEQS.update(b["nr"] for b in book)
    db = load_db(args.psql)
    updates, inserts, report = plan(book, db)

    Path(args.out).write_text(render_sql(updates, inserts, sum(b["netto"] for b in book)), encoding="utf-8")
    print("\n".join(report))
    print(f"\n{len(book)} Buchzeilen · {len(updates)} Änderungen · {len(inserts)} neue Zeilen · "
          f"{len({i['new_customer'] for i in inserts if i['new_customer']} | set(CUSTOMER_FROM_BOOK.values()))} neue Kunden")
    print(f"Buch netto gesamt: {eur(sum(b['netto'] for b in book))}")
    print(f"SQL: {args.out}")

    if args.apply:
        res = subprocess.run(args.psql + " -X -v ON_ERROR_STOP=1", shell=True,
                             input=Path(args.out).read_text(encoding="utf-8"), capture_output=True, text=True)
        print(res.stdout[-2000:])
        if res.returncode != 0:
            sys.exit(f"ABGEBROCHEN — nichts geschrieben:\n{res.stderr}")
        print("Angewendet.")


if __name__ == "__main__":
    main()
